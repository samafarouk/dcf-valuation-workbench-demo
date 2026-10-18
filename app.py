import io
from datetime import date

import pandas as pd
import streamlit as st

st.set_page_config(page_title="DCF Valuation Workbench", page_icon="📊", layout="wide")

NAVY = "#10233f"
BLUE = "#2476c7"

st.markdown(
    f"""
    <div style="background:linear-gradient(135deg,{NAVY},#174ea6);padding:24px 28px;border-radius:14px;color:white;margin-bottom:18px">
      <div style="font-size:12px;letter-spacing:2px;opacity:.85">FINANCIAL MODELING · EQUITY VALUATION</div>
      <div style="font-size:30px;font-weight:750;margin-top:6px">DCF Valuation Workbench</div>
      <div style="font-size:14px;opacity:.85;margin-top:5px">Forecast cash flows · Test scenarios · Explore valuation sensitivity</div>
    </div>
    """,
    unsafe_allow_html=True,
)

st.caption("Amounts are entered in EGP millions unless the field says otherwise. This is an analytical tool, not investment advice.")

REQUIRED_COLUMNS = ["Year", "EBIT", "Tax_Rate", "D&A", "CapEx", "Receivables", "Inventory", "Payables"]


def calculate_fcff(financial_data: pd.DataFrame) -> pd.DataFrame:
    df = financial_data.copy()
    df = df.sort_values("Year").reset_index(drop=True)
    df["NOPAT"] = df["EBIT"] * (1 - df["Tax_Rate"] / 100)
    df["NWC"] = df["Receivables"] + df["Inventory"] - df["Payables"]
    df["Change_NWC"] = df["NWC"].diff()
    df["FCFF"] = df["NOPAT"] + df["D&A"] - df["CapEx"] - df["Change_NWC"]
    return df


def calculate_dcf(current_fcff, growth_rate, wacc, terminal_growth, forecast_years, net_debt, shares_outstanding):
    if current_fcff <= 0:
        raise ValueError("Starting FCFF must be positive. Check the value or uploaded statements.")
    if growth_rate <= -1:
        raise ValueError("FCFF growth must be greater than -100%.")
    if wacc <= 0 or wacc <= terminal_growth:
        raise ValueError("WACC must be positive and greater than terminal growth.")
    if terminal_growth < 0:
        raise ValueError("Terminal growth cannot be negative.")
    if forecast_years < 1:
        raise ValueError("Forecast years must be at least 1.")
    if shares_outstanding <= 0:
        raise ValueError("Shares outstanding must be positive.")

    years = list(range(1, int(forecast_years) + 1))
    fcff_forecast = [current_fcff * (1 + growth_rate) ** year for year in years]
    discounted_fcff = [cf / (1 + wacc) ** year for year, cf in zip(years, fcff_forecast)]
    pv_cash_flows = sum(discounted_fcff)
    terminal_value = fcff_forecast[-1] * (1 + terminal_growth) / (wacc - terminal_growth)
    pv_terminal = terminal_value / (1 + wacc) ** forecast_years
    enterprise_value = pv_cash_flows + pv_terminal
    equity_value = enterprise_value - net_debt
    return {
        "Enterprise Value": enterprise_value,
        "Equity Value": equity_value,
        "Value per Share": equity_value / shares_outstanding,
        "PV of Forecast Cash Flows": pv_cash_flows,
        "Terminal Value": terminal_value,
        "PV of Terminal Value": pv_terminal,
        "FCFF Forecast": fcff_forecast,
        "Years": years,
    }


COLUMN_ALIASES = {
    "Year": ["year", "fy", "fiscal year", "period", "financial year", "date"],
    "EBIT": ["ebit", "operating profit", "operating income", "profit from operations", "operating earnings"],
    "Tax_Rate": ["tax_rate", "tax rate", "effective tax rate", "tax %", "tax rate %", "tax rate (%)"],
    "D&A": ["d&a", "da", "depreciation and amortization", "depreciation & amortization", "depreciation and amortisation", "depreciation & amortisation", "depreciation amortization", "depreciation", "depreciation and amortization expense"],
    "CapEx": ["capex", "capital expenditure", "capital expenditures", "purchase of property plant and equipment", "purchases of property plant and equipment", "additions to property plant and equipment", "capital spending"],
    "Receivables": ["receivables", "accounts receivable", "trade receivables", "trade and other receivables", "account receivables"],
    "Inventory": ["inventory", "inventories", "stocks"],
    "Payables": ["payables", "accounts payable", "trade payables", "trade and other payables", "account payables"],
}


def normalize_label(value):
    return " ".join(str(value).strip().lower().replace("_", " ").replace("&", " and ").replace("%", " percent").split())


def suggest_column_mapping(columns):
    normalized = {col: normalize_label(col) for col in columns}
    mapping = {}
    for target, aliases in COLUMN_ALIASES.items():
        choices = [normalize_label(target)] + [normalize_label(x) for x in aliases]
        match = next((col for col, norm in normalized.items() if norm in choices), None)
        if match is None:
            # Conservative partial matching for common labels; user confirms every selection.
            match = next((col for col, norm in normalized.items() if any(a in norm for a in choices if len(a) > 4)), None)
        mapping[target] = match if match is not None else "— Select a column —"
    return mapping


def read_uploaded_table(uploaded_file, sheet_name=None):
    uploaded_file.seek(0)
    if uploaded_file.name.lower().endswith(".csv"):
        return pd.read_csv(uploaded_file)
    return pd.read_excel(uploaded_file, sheet_name=sheet_name)


def convert_table_orientation(raw_data: pd.DataFrame, orientation: str) -> pd.DataFrame:
    """Convert a metric-by-year table into one row per year when requested."""
    data = raw_data.copy()
    data.columns = [str(c).strip() for c in data.columns]
    if orientation == "Each row is a year (standard)":
        return data

    if data.shape[1] < 2:
        raise ValueError("The horizontal layout needs a label column and at least one data column.")

    # In this layout, the first column contains metric labels (EBIT, D&A, etc.)
    # and the remaining columns contain years or periods.
    label_col = data.columns[0]
    data = data.dropna(subset=[label_col]).copy()
    data[label_col] = data[label_col].astype(str).str.strip()
    data = data[data[label_col] != ""]
    data = data.drop_duplicates(subset=[label_col], keep="first")
    data = data.set_index(label_col).T
    data.index.name = "Year"
    data = data.reset_index()
    return data


def validate_mapped_data(data):
    missing = [c for c in REQUIRED_COLUMNS if c not in data.columns]
    if missing:
        raise ValueError("Please map every required field before continuing: " + ", ".join(missing))
    data = data[REQUIRED_COLUMNS].copy()
    for col in REQUIRED_COLUMNS:
        data[col] = pd.to_numeric(data[col], errors="coerce")
    if data.empty:
        raise ValueError("The selected sheet has no data rows.")
    blank_cols = [c for c in REQUIRED_COLUMNS if data[c].isna().all()]
    if blank_cols:
        raise ValueError("No numeric values were found for: " + ", ".join(blank_cols) + ". Check your column mapping and selected sheet.")
    if data[REQUIRED_COLUMNS].isna().any().any():
        bad = [c for c in REQUIRED_COLUMNS if data[c].isna().any()]
        raise ValueError("Some cells are blank or non-numeric in: " + ", ".join(bad) + ". Fill or correct these values in the review table.")
    if data["Year"].duplicated().any():
        raise ValueError("Each year should appear only once. Check for duplicate periods.")
    if not data["Year"].mod(1).eq(0).all():
        raise ValueError("Year values should be whole numbers, such as 2024.")
    if ((data["Tax_Rate"] < 0) | (data["Tax_Rate"] > 100)).any():
        raise ValueError("Tax rate must be entered as a percentage from 0 to 100 (for example, 22 for 22%, not 0.22).")
    if len(data) < 2:
        raise ValueError("Upload at least two years of data so the model can calculate the change in working capital.")
    return data.sort_values("Year").reset_index(drop=True)


def make_excel_report(company, summary, forecast, scenarios, sensitivity, history=None):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Valuation Summary", index=False, startrow=3)
        forecast.to_excel(writer, sheet_name="FCFF Forecast", index=False, startrow=3)
        scenarios.to_excel(writer, sheet_name="Scenarios", index=False, startrow=3)
        sensitivity.to_excel(writer, sheet_name="Sensitivity", startrow=3)
        if history is not None:
            history.to_excel(writer, sheet_name="Historical FCFF", index=False, startrow=3)
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = writer.book
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = "A5"
            last_col = max(ws.max_column, 3)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
            ws.cell(1, 1).value = f"{company} | DCF Valuation Report"
            ws.cell(1, 1).font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
            ws.cell(1, 1).fill = PatternFill("solid", fgColor="10233F")
            ws.row_dimensions[1].height = 34
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
            ws.cell(2, 1).value = f"{ws.title} | Generated {date.today()} | Amounts in EGP millions unless stated otherwise"
            ws.cell(2, 1).font = Font(size=10, italic=True, color="64748B")
            for cell in ws[4]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2476C7")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[4].height = 30
            for col_cells in ws.columns:
                letter = get_column_letter(col_cells[0].column)
                max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=12)
                ws.column_dimensions[letter].width = min(max(max_len + 3, 14), 36)
    output.seek(0)
    return output.getvalue()

with st.sidebar:
    st.header("Valuation inputs")
    company = st.text_input("Company name", "Example Company")
    st.markdown("#### Cash flow assumptions")
    manual_fcff = st.number_input("Starting FCFF (EGP M)", min_value=0.01, value=100.0, step=5.0)
    growth_pct = st.number_input("Annual FCFF growth (%)", min_value=-99.0, max_value=100.0, value=8.0, step=1.0)
    wacc_pct = st.number_input("WACC (%)", min_value=0.1, max_value=100.0, value=12.0, step=0.5)
    terminal_pct = st.number_input("Terminal growth (%)", min_value=0.0, max_value=50.0, value=3.0, step=0.5)
    years = st.number_input("Forecast years", min_value=1, max_value=20, value=5, step=1)
    st.markdown("#### Capital structure")
    net_debt = st.number_input("Net debt (EGP M)", value=250.0, step=25.0)
    shares = st.number_input("Shares outstanding (millions)", min_value=0.01, value=100.0, step=10.0)
    calculate = st.button("Calculate valuation", type="primary", use_container_width=True)

st.subheader("1. Choose your FCFF source")
source = st.radio("How do you want to provide starting FCFF?", ["Enter FCFF manually", "Calculate FCFF from financial statements"], horizontal=True)
history = None
starting_fcff = manual_fcff
source_label = "Manual input"

if source == "Calculate FCFF from financial statements":
    st.markdown("Upload a structured Excel workbook or CSV. The app will suggest column matches, then let you confirm them before calculating FCFF.")
    with st.expander("What information is needed?", expanded=False):
        st.markdown("The model needs one row per financial year and eight inputs: Year, EBIT, Tax rate, D&A, CapEx, Receivables, Inventory, and Payables.")
        st.caption("Amounts should use one consistent unit and currency. Tax rate should be entered as 22 for 22%, not 0.22. For now, this guided importer reads structured tables; it does not automatically extract arbitrary PDF annual reports.")
    uploaded = st.file_uploader("Upload financial statements (.xlsx or .csv)", type=["xlsx", "csv"], key="financial_upload")
    if uploaded is not None:
        try:
            if uploaded.name.lower().endswith(".xlsx"):
                uploaded.seek(0)
                workbook = pd.ExcelFile(uploaded)
                sheet = st.selectbox("Which sheet contains the historical financial data?", workbook.sheet_names, key="financial_sheet")
                raw_upload = read_uploaded_table(uploaded, sheet)
            else:
                raw_upload = read_uploaded_table(uploaded)
            if raw_upload.empty or len(raw_upload.columns) == 0:
                st.error("The selected file or sheet appears to be empty.")
            else:
                raw_upload.columns = [str(c).strip() for c in raw_upload.columns]
                orientation = st.radio(
                    "How is your data arranged?",
                    ["Each row is a year (standard)", "Metrics are rows and years are columns (horizontal)"],
                    horizontal=False,
                    help="Choose standard if each row is a year and each column is a metric. Choose horizontal if the first column contains labels such as EBIT or D&A and the columns to the right contain years."
                )
                if orientation == "Metrics are rows and years are columns (horizontal)":
                    st.caption("We'll transpose this layout so each year becomes a row. The first column must contain metric names, and the remaining column headers should identify years.")
                    try:
                        raw_upload = convert_table_orientation(raw_upload, orientation)
                    except Exception as orientation_error:
                        st.error(f"We couldn't convert this layout: {orientation_error}")
                st.markdown("#### Step A — Preview your uploaded data")
                st.caption("Check that the correct table and header row were detected. If the preview starts with a title or blank rows, clean the sheet so the column names are in the first row of the table.")
                st.dataframe(raw_upload.head(12), use_container_width=True, hide_index=True)
                suggestions = suggest_column_mapping(raw_upload.columns)
                st.markdown("#### Step B — Confirm the column matches")
                st.caption("We suggest matches automatically. Please check them: similar names do not always mean the same financial concept.")
                options = ["— Select a column —"] + list(raw_upload.columns)
                mapped = {}
                left, right = st.columns(2)
                for i, target in enumerate(REQUIRED_COLUMNS):
                    container = left if i % 2 == 0 else right
                    with container:
                        suggested = suggestions.get(target, "— Select a column —")
                        index = options.index(suggested) if suggested in options else 0
                        mapped[target] = st.selectbox(
                            f"{target} " + ("(year)" if target == "Year" else "(percentage, 0–100)" if target == "Tax_Rate" else "(same currency/unit throughout)"),
                            options, index=index, key=f"map_{target}"
                        )
                unresolved = [target for target, col in mapped.items() if col == "— Select a column —"]
                duplicate_sources = [col for col in mapped.values() if col != "— Select a column —" and list(mapped.values()).count(col) > 1]
                if unresolved:
                    st.warning("Please choose a source column for: " + ", ".join(unresolved) + ".")
                if duplicate_sources:
                    st.warning("One source column is mapped to more than one field: " + ", ".join(sorted(set(duplicate_sources))) + ". Check that each mapping is correct.")
                if not unresolved:
                    candidate = pd.DataFrame({target: raw_upload[source_col] for target, source_col in mapped.items()})
                    st.markdown("#### Step C — Review the mapped inputs")
                    st.caption("You can edit incorrect or missing values in this review table before the model calculates FCFF.")
                    edited = st.data_editor(candidate, use_container_width=True, hide_index=True, num_rows="fixed", key="mapped_financial_review")
                    if st.button("Confirm data and calculate historical FCFF", type="primary", key="confirm_financial_data"):
                        try:
                            clean = validate_mapped_data(edited)
                            history = calculate_fcff(clean)
                            valid = history.dropna(subset=["FCFF"])
                            if valid.empty:
                                st.error("No valid FCFF could be calculated from the selected data.")
                                history = None
                            else:
                                latest = valid.iloc[-1]
                                starting_fcff = float(latest["FCFF"])
                                source_label = f"Uploaded statements — FY {int(latest['Year'])}"
                                st.session_state["confirmed_financial_history"] = history
                                st.session_state["confirmed_starting_fcff"] = starting_fcff
                                st.session_state["confirmed_source_label"] = source_label
                                st.success(f"Data confirmed. Latest calculated FCFF: {starting_fcff:,.2f} (FY {int(latest['Year'])}).")
                        except Exception as e:
                            st.error(f"Please review the mapped data: {e}")
                saved_history = st.session_state.get("confirmed_financial_history")
                if saved_history is not None:
                    history = saved_history
                    starting_fcff = float(st.session_state["confirmed_starting_fcff"])
                    source_label = st.session_state["confirmed_source_label"]
                    st.success(f"Using confirmed historical data. Latest FCFF: {starting_fcff:,.2f} — {source_label}.")
                    st.dataframe(history.round(2), use_container_width=True, hide_index=True)
        except Exception as e:
            st.error(f"We couldn't read this file or sheet: {e}")
            history = st.session_state.get("confirmed_financial_history")

if source == "Calculate FCFF from financial statements" and history is None:
    st.info("Upload a valid file first. The model will use the calculated latest FCFF after the upload is successful.")

if calculate:
    if not company.strip():
        st.error("Please enter a company name.")
    elif source == "Calculate FCFF from financial statements" and history is None:
        st.error("Please upload a valid financial-statement file before calculating.")
    else:
        try:
            g, r, tg = growth_pct / 100, wacc_pct / 100, terminal_pct / 100
            base = calculate_dcf(starting_fcff, g, r, tg, int(years), net_debt, shares)
            scenario_inputs = {
                "Downside": (g - 0.02, r + 0.02, max(0, tg - 0.01)),
                "Base": (g, r, tg),
            }
            up_r = max(r - 0.02, tg + 0.02)
            up_tg = min(tg + 0.01, up_r - 0.01)
            scenario_inputs["Upside"] = (g + 0.02, up_r, max(0, up_tg))
            scenario_rows = []
            for name, (sg, sr, stg) in scenario_inputs.items():
                if sr <= stg:
                    sr = stg + 0.01
                out = calculate_dcf(starting_fcff, sg, sr, stg, int(years), net_debt, shares)
                scenario_rows.append({"Scenario": name, "FCFF Growth": sg, "WACC": sr, "Terminal Growth": stg,
                                      "Enterprise Value (EGP M)": out["Enterprise Value"], "Equity Value (EGP M)": out["Equity Value"],
                                      "Value per Share (EGP)": out["Value per Share"]})
            scenarios = pd.DataFrame(scenario_rows)
            wacc_rates = sorted(set([max(0.001, r - 0.02), r, r + 0.02]))
            terminal_rates = sorted(set([max(0, tg - 0.01), tg, tg + 0.01]))
            sensitivity = pd.DataFrame(index=[f"{x:.1%}" for x in wacc_rates], columns=[f"{x:.1%}" for x in terminal_rates], dtype=float)
            for sr in wacc_rates:
                for stg in terminal_rates:
                    if sr <= stg:
                        sensitivity.loc[f"{sr:.1%}", f"{stg:.1%}"] = float("nan")
                    else:
                        out = calculate_dcf(starting_fcff, g, sr, stg, int(years), net_debt, shares)
                        sensitivity.loc[f"{sr:.1%}", f"{stg:.1%}"] = out["Enterprise Value"]
            forecast = pd.DataFrame({"Year": [f"Year {y}" for y in base["Years"]], "Projected FCFF (EGP M)": base["FCFF Forecast"]})
            summary = pd.DataFrame([{
                "Company": company, "Valuation Date": str(date.today()), "FCFF Source": source_label,
                "Starting FCFF (EGP M)": starting_fcff, "FCFF Growth": g, "WACC": r,
                "Terminal Growth": tg, "Forecast Years": int(years), "Net Debt (EGP M)": net_debt,
                "Shares Outstanding (M)": shares, "PV of Forecast Cash Flows (EGP M)": base["PV of Forecast Cash Flows"],
                "Terminal Value (EGP M)": base["Terminal Value"], "PV of Terminal Value (EGP M)": base["PV of Terminal Value"],
                "Enterprise Value (EGP M)": base["Enterprise Value"], "Equity Value (EGP M)": base["Equity Value"],
                "Value per Share (EGP)": base["Value per Share"]
            }])
            st.session_state["valuation_package"] = {"company": company, "summary": summary, "base": base,
                "scenarios": scenarios, "sensitivity": sensitivity, "forecast": forecast, "history": history}
        except Exception as e:
            st.error(f"Could not calculate valuation: {e}")

package = st.session_state.get("valuation_package")
if package:
    st.markdown("---")
    st.subheader(f"2. Valuation dashboard — {package['company']}")
    a, b, c = st.columns(3)
    a.metric("Enterprise value", f"EGP {package['base']['Enterprise Value']:,.2f}M")
    b.metric("Equity value", f"EGP {package['base']['Equity Value']:,.2f}M")
    c.metric("Intrinsic value per share", f"EGP {package['base']['Value per Share']:,.2f}")

    st.markdown("#### Forecast free cash flow")
    st.bar_chart(package["forecast"].set_index("Year"), y="Projected FCFF (EGP M)")
    st.dataframe(package["forecast"].round(2), use_container_width=True, hide_index=True)

    st.markdown("#### Downside / base / upside scenarios")
    st.dataframe(package["scenarios"].style.format({"FCFF Growth":"{:.1%}", "WACC":"{:.1%}", "Terminal Growth":"{:.1%}",
        "Enterprise Value (EGP M)":"{:,.2f}", "Equity Value (EGP M)":"{:,.2f}", "Value per Share (EGP)":"{:,.2f}"}), use_container_width=True, hide_index=True)

    st.markdown("#### Enterprise value sensitivity")
    st.caption("Rows are WACC; columns are terminal growth. N/A means WACC is not greater than terminal growth.")
    st.dataframe(package["sensitivity"].style.format("{:,.2f}", na_rep="N/A").background_gradient(cmap="RdYlGn", axis=None), use_container_width=True)

    with st.expander("Detailed valuation components"):
        base = package["base"]
        details = pd.DataFrame({"Component": ["PV of forecast cash flows", "Terminal value (at end of forecast)", "PV of terminal value"],
            "EGP millions": [base["PV of Forecast Cash Flows"], base["Terminal Value"], base["PV of Terminal Value"]]})
        st.dataframe(details.style.format({"EGP millions":"{:,.2f}"}), use_container_width=True, hide_index=True)

    st.markdown("#### Export your report")
    xlsx_bytes = make_excel_report(package["company"], package["summary"], package["forecast"], package["scenarios"], package["sensitivity"], package["history"])
    safe_name = "_".join(package["company"].strip().split()) or "Company"
    st.download_button("Download Excel valuation report", data=xlsx_bytes, file_name=f"{safe_name}_DCF_Valuation_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
    st.download_button("Download valuation summary (CSV)", data=package["summary"].to_csv(index=False).encode("utf-8"), file_name=f"{safe_name}_DCF_Summary.csv", mime="text/csv")
else:
    st.markdown("---")
    st.subheader("2. Your valuation results will appear here")
    st.write("Enter your assumptions in the left sidebar, choose your FCFF source, then click **Calculate valuation**.")

st.markdown("---")
st.caption("Model note: the valuation depends on the assumptions entered. Verify financial statement inputs, capital structure, and methodology before using results in investment decisions.")
